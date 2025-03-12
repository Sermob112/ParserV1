import csv
import os
import re
import json
import logging
from bs4 import BeautifulSoup
import pandas as pd
import requests
from selenium import webdriver
from selenium.webdriver.chrome.options import Options
from docx import Document
from ParserAgreement import AgreementParser
from selenium.webdriver.common.by import By
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from PySide6.QtCore import QObject, Signal
import time
# Настройка логгирования
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

class ParserOldR(QObject):
    message_signal = Signal(str)
    
    def __init__(self,file_name):
        super().__init__() 
        self.HEADERS = {
            'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/110.0.0.0 YaBrowser/23.3.0.2246 Yowser/2.5 Safari/537.36',
            'accept': '*/*'
        }
        self.status = None
        self.filePath = None
        self.main_directory = None
        self.ObjectName = None
        self.num = None
        self.file_name_to_version_controll = file_name
        self.soup = None
        self.contract_data_event = None
 

    def make_link_num(self, numer, filePath):
        self.filePath = filePath
        search_url = f'https://zakupki.gov.ru/epz/order/extendedsearch/results.html?searchString={numer}&morphology=on&search-filter=Дате+размещения&pageNumber=1&sortDirection=false&recordsPerPage=_10&showLotsInfoHidden=false&sortBy=UPDATE_DATE&fz44=on&fz223=on&af=on&ca=on&pc=on&pa=on&currencyIdGeneral=-1'
        
        try:
            req = requests.get(search_url, headers=self.HEADERS)
            req.raise_for_status()
            src = req.text
            soup = BeautifulSoup(src, 'lxml')
            col = soup.find('div', class_='registry-entry__header-mid__number')
            a_tag = col.find('a')
            match = re.search(r'noticeInfoId=(\d+)', a_tag['href'])
            link_text = 'https://zakupki.gov.ru/' + a_tag.get('href')
            number = match.group(1)
            self.main_directory = f'Закупка № {numer} '
            self.num = numer
            self.agent(number, link_text)
            
        except Exception as e:
           
            logger.error(f"Ошибка подключения: {e}")
            self.status = 'Ошибка подключения'

    def agent(self, numer, link_text):
        try:
            req = requests.get(link_text, headers=self.HEADERS)
            req.raise_for_status()
            src = req.text
            self.soup = BeautifulSoup(src, 'lxml')
            self.message_signal.emit(f"Успешное подключение: закупка № {numer}")
            self.status = 'Успешное подключение'
            return self.soup
        except Exception as e:
            self.message_signal.emit(f"Ошибка подключения: закупка № {numer}: {e}")
            logger.error(f"Ошибка подключения: {e}")
            self.status = 'Ошибка подключения'

    def parse_head(self):
        source = self.soup.find(class_="col-6 pr-0 mr-21px")
        self.ObjectName = source.find(class_='registry-entry__body-value').get_text().strip().replace('"', '').replace('\r', '')[:48].replace(' ', '_').replace('\n', '')
        return self.ObjectName

    def get_links(self):
        tabs_of_links = {}
        link_razdels = self.soup.find(class_='container card-layout')
        if link_razdels:
            try:
                link_razdels = self.soup.find(class_='tabsNav d-flex').find_all('a')
            except:
                link_razdels = self.soup.find(class_='tabsNav d-flex align-items-end').find_all('a')
            for links in link_razdels:
                linkl = 'https://zakupki.gov.ru/' + links.get('href')
                title_razde = links.text.strip()
                tabs_of_links[title_razde] = linkl
        return tabs_of_links

    def main_info(self):
        mainMass = []
        newmainMass = []
        JornalMass = []
        DockMass = []
        AgreMass = []
        tabs_of_links = self.get_links()
        for title, link in tabs_of_links.items():
            req = requests.get(url=link, headers=self.HEADERS)
            src = req.text
            soup = BeautifulSoup(src, "lxml")
            if title == 'Журнал событий':
                JornalMass = self.get_journals(link)
                self.version_controll(JornalMass,self.file_name_to_version_controll, self.contract_data_event)
            if title == 'Сведения о договорах':
                AgreMass = self.get_agreements(link)
            if title == 'Документы':
                DockMass = self.get_documents(link)
            try:
                containerMain = soup.find_all(class_='card-common-content')
                containerMain2 = soup.find_all(class_='card-attachments')
                if containerMain:
                    for i in containerMain:
                        lines = i.get_text().strip().splitlines()
                        Mass = [line.strip() for line in lines if line.strip()]
                        cleaned_Mass = [item.replace('\xa0', ' ') for item in Mass]
                        mainMass += cleaned_Mass
                if containerMain2:
                    for i in containerMain2:
                        lines = i.get_text().strip().splitlines()
                        Mass = [line.strip() for line in lines if line.strip()]
                        cleaned_Mass = [item.replace('\xa0', ' ') for item in Mass]
                        mainMass += cleaned_Mass
                        newmainMass += cleaned_Mass
            except Exception as e:
                logger.error(f"Ошибка при обработке контейнера: {e}")
                self.message_signal.emit(f"Ошибка при обработке контейнера: {e}")
        mainMass.append(newmainMass)
        mainMass.append(JornalMass)
        mainMass.append(DockMass)
        mainMass.append(AgreMass)
        return mainMass

    def get_agreements(self, link):
        req = requests.get(link, headers=self.HEADERS, params=None)
        src = req.text
        soup = BeautifulSoup(src, 'lxml')
        no_data_message = soup.find('div', class_='section__title', text='Сведения отсутствуют')
        if no_data_message:
            return ["Сведения отсутствуют"]
        table = soup.find('table', class_='table')
        if table:
            rows = table.find_all('tr')
            agreements = []
            for row in rows[1:]:
                cols = row.find_all('td')
                if len(cols) >= 4:
                    links = cols[0].find_all('a', href=True)
                    self.contract_info_link = None
                    for link in links:
                        if 'contract-info' in link['href']:
                            self.contract_info_link = link['href']
                    agreement_text = cols[0].get_text(strip=True)
                    agreement_revision = cols[1].get_text(strip=True)
                    agreement_status = cols[2].get_text(strip=True)
                    agreement_date = cols[3].get_text(strip=True)
                    agreements.append(f"Текст: {agreement_text}")
                    agreements.append(f"Редакция: {agreement_revision}")
                    agreements.append(f"Статус: {agreement_status}")
                    agreements.append(f"Дата: {agreement_date}")
                    agreements.append("\n")
            try:
                link = "https://zakupki.gov.ru" + self.contract_info_link
                par_agreement = AgreementParser(link)
                par_agreement.message_signal.connect(self.message_signal.emit)
                par_agreement.start_contract_parser(file_path=f"{self.filePath}/{self.main_directory + self.ObjectName}")
                self.contract_data_event = par_agreement.get_event_data()
            except Exception as e:
                self.message_signal.emit(f"Ошибка при парсинге договора: {e}")
                logger.error(f"Ошибка при парсинге договора: {e}")
            return agreements
        else:
            return ["Таблица с договорами не найдена"]
        
    
    def get_journals(self, link):
        chrome_options = Options()
        chrome_options.add_argument("--headless")  # Запуск без GUI
        chrome_options.add_argument("--disable-gpu")
        chrome_options.add_argument("--window-size=1920,1080")
        
        driver = webdriver.Chrome(options=chrome_options)
        driver.get(link)

        all_data = []  # <-- Инициализация переменной до блока try

        try:
            # Закрываем модальное окно, если оно есть
            try:
                close_button = WebDriverWait(driver, 5).until(
                    EC.element_to_be_clickable((By.CLASS_NAME, 'btn-close'))
                )
                close_button.click()
                time.sleep(1)  # Даем время на закрытие
            except:
                print("Модального окна нет или уже закрыто.")

            # Проверяем, есть ли выпадающий список для выбора количества записей на странице
            select_boxes = driver.find_elements(By.CLASS_NAME, 'select-record-per-page--number')
            if select_boxes:
                select_box = select_boxes[0]  # Берем первый найденный элемент
                select_box.click()
                time.sleep(1)

                # Проверяем, есть ли опция "50"
                option_50s = driver.find_elements(By.ID, '_50')
                if option_50s:
                    option_50 = option_50s[0]
                    option_50.click()
                    time.sleep(3)  # Даем время на подгрузку данных
                else:
                    print("Опция '50' не найдена, продолжаем с дефолтным значением.")
            else:
                print("Выпадающий список пагинации не найден, продолжаем.")

            # Теперь парсим данные
            WebDriverWait(driver, 10).until(
                EC.presence_of_element_located((By.CLASS_NAME, 'tabBoxWrapper'))
            )

            html = driver.page_source
            soup = BeautifulSoup(html, 'lxml')

            containerMain3 = soup.find(class_='tabBoxWrapper tabBoxWrapper__mb24')
            if containerMain3:
                headers = [th.get_text().strip().replace('\n', '') for th in containerMain3.find_all('th')]
                rows = [td.get_text().strip().replace('\n', '') for td in containerMain3.find_all('td')]
                all_data = [dict(zip(headers, rows[i:i + len(headers)])) for i in range(0, len(rows), len(headers))]

        except Exception as e:
            self.message_signal.emit(f"Ошибка при парсинге Журнал Событий: {e}")
            print("Ошибка в методе get_journals:", e)

        driver.quit()
        return all_data    
    def contrac_event_data_rework(self, data):
        result = []
        
        # Проверяем, есть ли ключ 'Журнал событий' и 'События', а также является ли 'События' списком
        if 'Журнал событий' in data and 'События' in data['Журнал событий'] and isinstance(data['Журнал событий']['События'], list):
            for event in data['Журнал событий']['События']:  # Проходим по каждому событию
                result.append({
                    'Дата и время': event.get('Дата и время', ''),
                    'Событие': event.get('Событие', '')
                })
        
        return result
        
    
    def version_controll(self, data, file_path,contrac_event_data_rework):
        if contrac_event_data_rework != None:
            combained_data = data + self.contrac_event_data_rework(contrac_event_data_rework)
        else: 
            combained_data = data
      
        try:
       
            try:
                df = pd.read_csv(file_path, encoding='windows-1251', sep=';', dtype=str)
            except FileNotFoundError:
                df = pd.DataFrame()

            df.columns = df.columns.str.strip()
            if "Реестровый номер закупки" not in df.columns:
                print("В файле нет нужного столбца 'Реестровый номер закупки'.")
                return

         
            df["Реестровый номер закупки"] = df["Реестровый номер закупки"].astype(str).str.replace("№", "")

            
            matching_rows = df[df["Реестровый номер закупки"] == self.num]

            if not matching_rows.empty:
                row_index = matching_rows.index[0]  # Берем индекс первой найденной строки

                
                new_row = {}
                for i, entry in enumerate(combained_data):
                    new_row[f'Дата и время_{i+1}'] = entry['Дата и время']
                    new_row[f'Событие_{i+1}'] = entry['Событие']

                
                for key, value in new_row.items():
                    df.at[row_index, key] = value

                print(f"Данные обновлены в строке с реестровым номером {self.num}.")
            else:
                print(f"Реестровый номер {self.num} не найден в файле.")

            df["Реестровый номер закупки"] = "№" + df["Реестровый номер закупки"]

            df.to_csv(file_path, index=False, encoding='windows-1251', sep=';', 
                    lineterminator='\n', quoting=csv.QUOTE_NONNUMERIC)

            self.status = 'Успешное обновление Журнала'
        except Exception as e:
            self.message_signal.emit(f"Ошибка обработки Журнала событий: {e}")
            print(f'Ошибка обработки Журнала событий: {e}')
    def get_documents(self, link):
        titlesMass = []
        linkMass = []
        req = requests.get(link, headers=self.HEADERS, params=None)
        src = req.text
        soup = BeautifulSoup(src, 'lxml')
        findTitle = soup.find_all(class_='container card-attachments-container')
        for titles in findTitle:
            title = titles.find(class_='title')
            if title:
                title_text = title.get_text().strip()
                values = titles.find_all(class_='col-6 b-left')
                for laxir in values:
                    luxit = laxir.find_all('a')
                    try:
                        for links in luxit:
                            href = links.get('href')
                            if href and 'download/download.html?id' in href:
                                linkl = href
                                tooltip_value = links.get('data-tooltip')
                                tooltip_soup = BeautifulSoup(tooltip_value, 'html.parser')
                                span_element = tooltip_soup.find('span')
                                if span_element:
                                    text = span_element.get_text(strip=True).replace('"', '')
                                    response = requests.get(f'https://zakupki.gov.ru{linkl}', headers=self.HEADERS)
                                    if response.status_code == 200:
                                        os.makedirs(f'{self.filePath}/{self.main_directory + self.ObjectName}/{title_text}', exist_ok=True)
                                        with open(f'{self.filePath}/{self.main_directory + self.ObjectName}/{title_text}/{text}', "wb") as f:
                                            f.write(response.content)
                                    else:
                                        logger.error(f"Не удалось скачать файл: {linkl}")
                                        self.message_signal.emit(f"Не удалось скачать файл: {linkl}")
                    except Exception as e:
                        logger.error(f"Произошла ошибка: {e}")
                        self.message_signal.emit(f"Ошибка при парсинге Документов: {e}")
    
    def make_doc(self):
        mainMass = self.main_info()

        try:
            doc = Document()
            doc.add_heading('Данные о закупке', level=1)
            for item in mainMass:
                if isinstance(item, dict):  # Если item — словарь, форматируем его
                    for key, value in item.items():
                        doc.add_paragraph(f"{key}: {value}")
                else:
                    doc.add_paragraph(str(item))
            doc.save(f"{self.filePath}/{self.main_directory + self.ObjectName}/Все данные о закупке №{self.num}.docx")
            self.message_signal.emit(f"Успешная запись файлов №{self.num}")
        except Exception as e:
            self.message_signal.emit(f"Ошибка записи файлов: {e}")
            logger.error(f"Ошибка записи файлов: {e}")
            self.status = 'Ошибка записи файлов'

# if __name__ == "__main__":
#         parR= ParserOldR()
#         parR.parse_head()
#         parR.documents(i)
#         parR.get_supplier_links(i)
#         parR.get_result_contracts(i)
#         parR.Make_Dock(i)
