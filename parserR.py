import csv
import pandas as pd
from bs4 import BeautifulSoup
import requests
import re
from typing import Dict, List, Optional
import os
from selenium import webdriver
from selenium.webdriver.chrome.options import Options
from docx import Document
from contract_parser import ContractParser
from PySide6.QtCore import QObject, Signal
from selenium.webdriver.common.by import By
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from urllib.parse import unquote
import time

class ParserR(QObject):
    message_signal = Signal(str)
  
    def __init__(self,file_name):
        super().__init__() 
        self.log_data = []
        self.headers = {
            'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/110.0.0.0 YaBrowser/23.3.0.2246 Yowser/2.5 Safari/537.36',
            'accept': '*/*'
        }
        self.status = None
        self.soup = None
        self.main_directory = None
        self.object_name = None
        self.contract_data_event = None
        self.file_name_to_version_controll = file_name
        self.contract_parser = None
        
        

    def make_link_num(self, numer: str, file_path: str) -> None:
        self.file_path = file_path
        search_url = f'https://zakupki.gov.ru/epz/order/extendedsearch/results.html?searchString={numer}&morphology=on&search-filter=Дате+размещения&pageNumber=1&sortDirection=false&recordsPerPage=_10&showLotsInfoHidden=false&sortBy=UPDATE_DATE&fz44=on&fz223=on&af=on&ca=on&pc=on&pa=on&currencyIdGeneral=-1'
        
        try:
            req = requests.get(search_url, headers=self.headers)
            req.raise_for_status()
            src = req.text
            soup = BeautifulSoup(src, 'lxml')
            col = soup.find('div', class_='registry-entry__header-mid__number')
            a_tag = col.find('a')
            link_text = 'https://zakupki.gov.ru/' + a_tag.get('href')
            self.main_directory = f'Закупка № {numer}'
            self.num = numer
            self.agent(file_path, link_text, numer)
            self.status = 'Успешное подключение'
            
        except Exception as e:
            
            self.status = f'Ошибка подключения: {e}'

    def agent(self, file_path: str, link: str,numer) -> Optional[BeautifulSoup]:
        self.file_path = file_path
        try:
            req = requests.get(link, headers=self.headers)
            req.raise_for_status()
            src = req.text
            self.soup = BeautifulSoup(src, 'lxml')
            self.status = 'Успешное подключение'
            self.message_signal.emit(f"Закупка № {numer} успешное подключение")
            return self.soup
        except Exception as e:
            self.status = f'Ошибка подключения: {e}'
            self.message_signal.emit(f"Ошибка подключения: закупка № {numer}: {e}")
            return None

    def parse_head(self) -> Dict:
        mass = []
        serial_date = {}
        
        try:
            serial_number = self.soup.find(class_="registry-entry__header-mid__number")
            name = self.soup.find(class_='cardMainInfo__content').get_text().strip().replace('"', '').replace('\r', '')[:60].replace(' ', '_').replace('\n', '')
            self.object_name = f' {name}'
            
            if serial_number:
                serial_number = serial_number.find('a').text.strip()
            else:
                serial_number = self.soup.find('a', {'class': 'distancedText'})
                reg_number = serial_number['href'].split('=')[1]
                status_text = self.soup.find(class_='cardMainInfo__state distancedText').get_text().strip()
                mass.append({reg_number: status_text})
                self.main_directory = f'Закупка № {reg_number}'

            svedenia_o_zakupke = self._parse_zakupchik()
            mass.append(svedenia_o_zakupke)

            price_info = self._parse_price()
            for i in range(0, len(price_info), 2):
                key = price_info[i]
                value = price_info[i+1] if i+1 < len(price_info) else ''
                mass.append({key: value})

            serial_date['Номер заказа'] = mass
            self.status = 'Успешный парсинг заголовка'
            return serial_date
        except Exception as e:
            self.status = f'Произошла ошибка с парсингом заголовка: {e}'
            self.message_signal.emit(f"Произошла ошибка с парсингом заголовка: {e}")
            return {}

    def _parse_zakupchik(self) -> List[Dict]:
        svedenia_o_zakupke = []
        zakupchik = self.soup.find(class_='registry-entry__body') or self.soup.find(class_='sectionMainInfo__body')
        if zakupchik:
            blocks = zakupchik.find_all(class_='registry-entry__body-block') or zakupchik.find_all(class_='cardMainInfo__section')
            for block in blocks:
                title = block.find(class_='registry-entry__body-title') or block.find(class_='cardMainInfo__title')
                value = block.find(class_='registry-entry__body-value') or block.find(class_='cardMainInfo__content')
                if title and value:
                    svedenia_o_zakupke.append({title.text.strip(): value.text.strip()})
                a_tag = block.find('a')
                if a_tag:
                    link = 'https://zakupki.gov.ru' + a_tag["href"]
                    svedenia_o_zakupke.append({'ссылка на организацию': link})
        return svedenia_o_zakupke

    def _parse_price(self) -> List[str]:
        price_info = []
        price = self.soup.find(class_='price-block') or self.soup.find(class_='sectionMainInfo borderRight col-3 colSpaceBetween')
        if price:
            elements = price.find_all('div') or price.find_all('span')
            for element in elements:
                title = element.text.strip()
                if '\xa0' in title:
                    title = title.replace('\xa0', '').replace(',', '.')
                price_info.append(title)
        return price_info
    
    def all_link(self) -> Dict[str, str]:
        """
        Собирает все ссылки из раздела "tabsNav" и возвращает их в виде словаря.
        """
        tabs_of_links = {}
        link_razdels = self.soup.find(class_='tabsNav d-flex align-items-end') or self.soup.find(class_='tabsNav d-flex')
        
        if link_razdels:
            try:
                links = link_razdels.find_all('a')
                for link in links:
                    link_url = 'https://zakupki.gov.ru/' + link.get('href')
                    title = link.text.strip()
                    tabs_of_links[title] = link_url
            except Exception as e:
                self.status = f'Ошибка при парсинге ссылок: {e}'
                self.message_signal.emit(f"Ошибка при парсинге ссылок: {e}")
        
        return tabs_of_links

    def main_info_body(self, soup: BeautifulSoup) -> Dict[str, List[Dict]]:
        """
        Парсит основную информацию о закупке и возвращает её в структурированном виде.
        """
        new_dict = {}
        try:
            main_infp = soup.find_all(class_='common-text b-bottom pb-3') or soup.find_all(class_='row blockInfo')
            
            if not main_infp:
                self.status = 'Нет данных для парсинга'
                return new_dict

            for info in main_infp:
                list_info = []
                title_cap = self._get_title(info)
                if not title_cap:
                    continue

                if 'common-text b-bottom pb-3' in info.get('class', []):
                    list_info.extend(self._parse_common_text(info))
                else:
                    list_info.extend(self._parse_block_info(info))

                if list_info:
                    new_dict[title_cap] = list_info

            self.status = 'Успешный парсинг общей информации'
            return new_dict
        except Exception as e:
            self.status = f'Ошибка при парсинге общей информации: {e}'
            self.message_signal.emit(f"Ошибка при парсинге общей информации: {e}")
            return new_dict

    def _get_title(self, info) -> Optional[str]:
        """
        Извлекает заголовок блока информации.
        """
        title = info.find(class_='common-text__caption') or info.find(class_='blockInfo__title')
        return title.text.strip() if title else None

    def _parse_common_text(self, info) -> List[Dict]:
        """
        Парсит блоки с классом 'common-text b-bottom pb-3'.
        """
        list_info = []
        col_9 = info.find_all(class_='col-9 mr-auto')
        
        for values in col_9:
            title_text = values.find(class_='common-text__title')
            value_text = values.find(class_='common-text__value')
            
            if title_text and value_text:
                title_text = title_text.text.strip()
                value_text = value_text.text.strip()
                list_info.append({title_text: value_text})
        
        return list_info

    def _parse_block_info(self, info) -> List[Dict]:
        """
        Парсит блоки с классом 'row blockInfo'.
        """
        list_info = []
        col_9 = info.find_all(class_='blockInfo__section')
        
        for values in col_9:
            title_text = values.find(class_='section__title')
            value_text = values.find(class_='section__info')
            
            if title_text and value_text:
                title_text = title_text.get_text().strip()
                value_text = value_text.get_text().strip()
                title_text = title_text.replace('\n', '')
                value_text = value_text.replace('\n', '').replace(' ', '')
                list_info.append({title_text: value_text})

        table = info.find(class_='blockInfo__table tableBlock')
        if table:
            list_info.extend(self._parse_table(table))
        
        return list_info

    def _parse_table(self, table) -> List[Dict]:
        """
        Парсит таблицу в блоке информации.
        """
        list_info = []
        headers = [th.text.strip() for th in table.find_all('th')]
        rows = table.find_all('td')
        
        for idx, row in enumerate(rows):
            td = row.get_text().strip()
            if not td:
                continue
            
            td = td.replace('\n', '').replace('\xa0', '').replace(',', '.')
            if re.match(r'[A-ZА-Я]', td):
                td = re.findall(r'[A-ZА-Я][^A-ZА-Я]*', td)
                td = ' '.join(td)
            
            header_idx = idx % len(headers)
            list_info.append({headers[header_idx]: td})
        
        return list_info
    def get_text(self,element) -> str:
        """Извлекает текст из элемента и очищает его."""
        return element.get_text().replace('\n', '').replace('\xa0', ' ').strip()

    def parse_collapse_title(self,collapse_element) -> str:
        """Извлекает заголовок collapse элемента."""
        title_element = collapse_element.find(class_='collapse__title_text')
        return self.get_text(title_element) if title_element else ''
    
    def parse_block_info(self,block_info_element) -> Dict[str, str]:
        """Парсит блок информации и возвращает словарь с данными."""
        data = {}
        title_element = block_info_element.find(class_='blockInfo__title')
        if title_element:
            title = self.get_text(title_element)
            sections = block_info_element.find_all(class_='section__title')
            infos = block_info_element.find_all(class_='section__info')
            for section, info in zip(sections, infos):
                section_title = self.get_text(section)
                data[section_title] = self.get_text(info)
        return data
    def parse_table(self,table_element) -> Dict[str, List[Dict[str, str]]]:
        """Парсит таблицу и возвращает словарь с данными."""
        table_data = {}
        headers = [self.get_text(th) for th in table_element.find_all('th')]
        rows = table_element.find_all('tr')
        for row in rows:
            cells = row.find_all('td')
            if cells:
                row_data = {}
                for header, cell in zip(headers, cells):
                    row_data[header] = self.get_text(cell)
                table_data.setdefault(headers[0], []).append(row_data)
        return table_data
    def collapse_element(self, soup: BeautifulSoup) -> Dict[str, Dict]:
        """Основная функция для парсинга collapse элементов."""
        collaps_data = {}
        collaps_content = soup.find_all('div', class_='blockInfo__collapse collapseInfo')

        for collapse in collaps_content:
            title = self.parse_collapse_title(collapse)
            if not title:
                continue

            collaps_data[title] = {}
            content_blocks = collapse.find_all('div', class_='content__block blockInfo')
            for block in content_blocks:
                block_data = self.parse_block_info(block)
                collaps_data[title].update(block_data)

            tables = collapse.find_all('table')
            for table in tables:
                table_data = self.parse_table(table)
                collaps_data[title].update(table_data)

        self.status = 'успешный парсинг выпадающих элементов'
        return collaps_data
    def documents(self, num: str) -> Dict[str, List]:
        """
        Парсит документы, связанные с закупкой, и сохраняет их на диск.
        Возвращает структурированные данные о документах.
        """
        data = {}
        link = f'https://zakupki.gov.ru/epz/order/notice/ea20/view/documents.html?regNumber={num}'
        
        try:
            req = requests.get(url=link, headers=self.headers)
            req.raise_for_status()
            soup = BeautifulSoup(req.text, "lxml")
            col_sm_12 = soup.find_all(class_='col-sm-12 blockInfo')
            
            for col in col_sm_12:
                titles = col.find(class_='blockInfo__title').get_text().strip()
                infos = self._parse_document_section(col)
                
                if infos:
                    data[titles] = infos
                    self._download_files(col, titles)
                
            self.status = 'Успешный парсинг документов'
        except Exception as e:
            self.status = f'Ошибка парсинга документов: {e}'
        
        return data

    def _parse_document_section(self, col) -> List:
        """
        Парсит секцию документа и возвращает структурированные данные.
        """
        infos = []
        section_value = col.find_all(class_='section__value docName')
        
        for sec in section_value:
            infos.append(sec.get_text().strip())
        
        sec_values = col.find_all(class_='section__value')
        for val in sec_values:
            infos.append(val.get_text().strip())
        
        sec_attribs = col.find_all(class_='section__attrib')
        for attrib in sec_attribs:
            infos.append(attrib.get_text().strip().replace('\n', ''))
        
        front = col.find_all('div', attrs={'style': 'font-size: 14px'})
        for f in front:
            infos.append(f.get_text().strip())
        
        return infos

    def _download_files(self, col, titles: str) -> None:
        """
        Скачивает файлы, связанные с документом, и сохраняет их на диск.
        """
        files_links = {}
        link_of_files = col.find_all(class_='blockFilesTabDocs')
        
        for lux in link_of_files:
            luxit = lux.find_all('a')
            for links in luxit:
                if 'download' in links.get('href') or 'file' in links.get('href'):
                    linkl = links.get('href')
                    titk = links.get('title')
                    files_links[titk] = linkl
        
        if files_links:
            dir_path = os.path.join(self.file_path, self.main_directory + self.object_name, titles)
            os.makedirs(dir_path, exist_ok=True)
            
            for title, url in files_links.items():
                try:
                    response = requests.get(url, headers=self.headers)
                    response.raise_for_status()
                    with open(os.path.join(dir_path, title), "wb") as f:
                        f.write(response.content)
                except Exception as e:
                    print(f'Ошибка при скачивании файла {title}: {e}')


    def _get_page_source(self, link):
        """Запускает браузер в headless-режиме и возвращает HTML-страницу."""
        chrome_options = Options()
        chrome_options.add_argument("--headless")
        chrome_options.add_argument("--disable-gpu")
        chrome_options.add_argument("--window-size=1920,1080")
        
        driver = webdriver.Chrome(options=chrome_options)
        driver.get(link)
        page_source = driver.page_source
        driver.quit()
        
        return page_source
    def get_supplier_links(self, num):
        """Получает ссылки на документы поставщика."""
        link_mass = []
        link = f'https://zakupki.gov.ru/epz/order/notice/ok20/view/supplier-results.html?regNumber={num}'
        
        try:
            response = requests.get(url=link, headers=self.headers)
            soup = BeautifulSoup(response.text, "lxml")
            
            for row in soup.find_all(class_='row blockInfo'):
                for a_tag in row.find_all('a', href=True):
                    full_link = 'https://zakupki.gov.ru' + a_tag['href']
                    
                    if 'protocol-main-info.html' in full_link:
                        full_link = full_link.replace('protocol-main-info.html', 'protocol-docs.html')
                        self.get_supplier_docs(full_link)
                    elif 'contractCard' in full_link:
                        self.contract_parser = ContractParser(full_link)
                        self.contract_parser.message_signal.connect(self.message_signal.emit)
                        self.contract_parser.start(f"{self.file_path}/{self.main_directory + self.object_name}")
                       
                        self.contract_data_event = self.contract_parser.get_event_data()
                        full_link = full_link.replace('common-info.html', 'document-info.html')
                        self.get_contract_details(full_link)
                    elif 'rdik/card/info.html' in full_link:
                        self.get_electronic_documents(full_link)
                    
                    link_mass.append(full_link)
        except Exception as e:
            print(f"Ошибка получения ссылок: {e}")
        
        return link_mass
    # def journal_of_events(self, link):
    #     """Парсит журнал событий."""
    #     try:
    #         soup = BeautifulSoup(self._get_page_source(link), 'lxml')
    #         table = soup.find(class_='table mb-0 displaytagTable')
            
    #         if not table:
    #             self.status = 'Журнал не найден'
    #             return []
            
    #         headers = [th.get_text().strip() for th in table.find_all('th')]
    #         rows = [td.get_text().strip() for td in table.find_all('td')]
            
    #         data = [dict(zip(headers, rows[i:i+len(headers)])) for i in range(0, len(rows), len(headers))]
            
    #         self.status = 'Успешный парсинг Журнала'
    #         # print(data)
    #         return data
    #     except Exception as e:
    #         self.status = f'Ошибка парсинга Журнала: {e}'
    #         self.message_signal.emit(f"Ошибка парсинга Журнала: {e}")
    #         return []
        
    def journal_of_events(self, link):
        chrome_options = Options()
        chrome_options.add_argument("--headless")  # Запуск без GUI
        chrome_options.add_argument("--disable-gpu")
        chrome_options.add_argument("--window-size=1920,1080")
        
        driver = webdriver.Chrome(options=chrome_options)
        driver.get(link)

        all_data = []  # <-- Инициализация переменной до блока try

        try:
            # Закрываем модальное окно, если оно есть
            try:
                close_button = WebDriverWait(driver, 5).until(
                    EC.element_to_be_clickable((By.CLASS_NAME, 'btn-close'))
                )
                close_button.click()
                time.sleep(1)  # Даем время на закрытие
            except:
                print("Модального окна нет или уже закрыто.")

            # Проверяем, есть ли выпадающий список для выбора количества записей на странице
            select_boxes = driver.find_elements(By.CLASS_NAME, 'select-record-per-page--number')
            if select_boxes:
                select_box = select_boxes[0]  # Берем первый найденный элемент
                select_box.click()
                time.sleep(1)

                # Проверяем, есть ли опция "50"
                option_50s = driver.find_elements(By.ID, '_50')
                if option_50s:
                    option_50 = option_50s[0]
                    option_50.click()
                    time.sleep(3)  # Даем время на подгрузку данных
                else:
                    print("Опция '50' не найдена, продолжаем с дефолтным значением.")
            else:
                print("Выпадающий список пагинации не найден, продолжаем.")

            # Теперь парсим данные
            WebDriverWait(driver, 10).until(
                EC.presence_of_element_located((By.CLASS_NAME, 'tabBoxWrapper'))
            )

            html = driver.page_source
            soup = BeautifulSoup(html, 'lxml')

            containerMain3 = soup.find(class_='table mb-0 displaytagTable')
            if containerMain3:
                headers = [th.get_text().strip().replace('\n', '') for th in containerMain3.find_all('th')]
                rows = [td.get_text().strip().replace('\n', '') for td in containerMain3.find_all('td')]
                all_data = [dict(zip(headers, rows[i:i + len(headers)])) for i in range(0, len(rows), len(headers))]

        except Exception as e:
            self.message_signal.emit(f"Ошибка при парсинге Журнал Событий\Версий: {e}")
            print("Ошибка в методе get_journals:", e)

        driver.quit()
        return all_data    
        
    def version_controll(self, data, file_path, contrac_event_data_rework):
        if contrac_event_data_rework != None:
            combained_data = data + self.contrac_event_data_rework(contrac_event_data_rework)
            print("Списки соеденены")
            # print(combained_data)
        else: 
            combained_data = data
            print("Списки не соеденены")
      
        try:
            # Читаем CSV, загружаем все как строки, чтобы избежать преобразования номеров
            try:
                df = pd.read_csv(file_path, encoding='windows-1251', sep=';', dtype=str)
            except FileNotFoundError:
                df = pd.DataFrame()
            df.columns = df.columns.str.strip()
            # Проверяем наличие столбца
            if "Реестровый номер закупки" not in df.columns:
                print("В файле нет нужного столбца 'Реестровый номер закупки'.")
                return

            # Убираем "№" для поиска
            df["Реестровый номер закупки"] = df["Реестровый номер закупки"].astype(str).str.replace("№", "")

            # Ищем строку с соответствующим реестровым номером
            matching_rows = df[df["Реестровый номер закупки"] == self.num]

            if not matching_rows.empty:
                row_index = matching_rows.index[0]  # Берем индекс первой найденной строки

                # Формируем новые данные
                new_row = {}
                for i, entry in enumerate(combained_data):
                    new_row[f'Дата и время_{i+1}'] = entry['Дата и время']
                    new_row[f'Событие_{i+1}'] = entry['Событие']

                # Обновляем строку
                for key, value in new_row.items():
                    df.at[row_index, key] = value

                print(f"Данные обновлены в строке с реестровым номером {self.num}.")
            else:
                print(f"Реестровый номер {self.num} не найден в файле.")

            # ВОЗВРАЩАЕМ "№" перед сохранением
            df["Реестровый номер закупки"] = "№" + df["Реестровый номер закупки"]

            # Сохраняем CSV без преобразования в экспоненциальную запись
            df.to_csv(file_path, index=False, encoding='windows-1251', sep=';', 
                    lineterminator='\n', quoting=csv.QUOTE_NONNUMERIC)

            self.status = 'Успешное обновление Журнала'
            self.message_signal.emit(f"Версии закупки {self.num} и контракта успешно сохранен")
        except Exception as e:
            self.message_signal.emit(f"Ошибка обработки Журнала событий: {e}")
            print(f'Ошибка обработки Журнала событий: {e}')
                    
         

    def supplier_result(self, link):
        """Парсит информацию о поставщиках."""
        try:
            soup = BeautifulSoup(self._get_page_source(link), 'lxml')
            sections = soup.find_all(class_='row blockInfo')
            data = {}
            
            for section in sections:
                title = section.find(class_='blockInfo__title').get_text(strip=True)
                info = {}
                
                for inf in section.find_all('section', class_='blockInfo__section section'):
                    key = inf.find(class_='section__title').get_text(strip=True)
                    value = inf.find(class_='section__info')
                    info[key] = value.get_text(strip=True) if value else ''
                
                data[title] = info
            
            self.status = 'Успешный парсинг поставщиков'
            return data
        except Exception as e:
            self.message_signal.emit(f"Ошибка парсинга поставщиков: {e}")
            self.status = f'Ошибка парсинга поставщиков: {e}'
            return {}
    def other_info(self):
        """Сбор дополнительной информации."""
        data = {}
        linkers = self.all_link()
        data_journal = {}
        for title, link in linkers.items():
            try:
                response = requests.get(url=link, headers=self.headers)
                soup = BeautifulSoup(response.text, "lxml")
                
                if self.main_info_body(soup):
                    data[title] = self.main_info_body(soup)
                elif title == 'Документы':
                    data[title] = self.documents(soup)
                elif title == 'Результаты определения поставщика (подрядчика, исполнителя)':
                    data[title] = self.supplier_result(link)
                    self.get_supplier_docs()
                elif title == 'Журнал событий':
                    data_journal = self.journal_of_events(link)
                    data[title] = data_journal
                    # print(self.file_name_to_version_controll)
                    self.version_controll(data_journal, self.file_name_to_version_controll, self.contract_data_event)
                   
            except Exception as e:
                self.message_signal.emit(f"Ошибка обработки {title}: {e}")
                print(f"Ошибка обработки {title}: {e}")
        
        return data

    
    
    def contrac_event_data_rework(self, data):
        result = []
    
        # Проверяем, есть ли ключ 'Динамические таблицы' и не пуст ли он
        if 'Динамические таблицы' in data and isinstance(data['Динамические таблицы'], list):
            for table in data['Динамические таблицы']:
                for event in table:  # Проходим по каждому словарю внутри списка
                    result.append({
                        'Дата и время': event.get('Дата и время', ''),
                        'Событие': event.get('Событие', '')
                    })
        
        return result

    
    def get_supplier_docs(self, link):
        """Получает и скачивает документы поставщика."""
        try:
            response = requests.get(url=link, headers=self.headers)
            soup = BeautifulSoup(response.text, "lxml")
            
            for block in soup.find_all(class_='row blockInfo'):
                title = block.find(class_='blockInfo__title')
                if title:
                    title_text = title.get_text().strip()
                    
                    for link_tag in block.find_all('a', href=True):
                        file_link = link_tag['href']
                        file_name = link_tag.get_text().strip()
                        
                        if 'download' in file_link or 'file' in file_link:
                            self.downloader(file_link, title_text, file_name)
        except Exception as e:
            print(f"Ошибка при получении документов: {e}")
            self.message_signal.emit(f"Ошибка при получении документов: {e}")

    def downloader(self, link, title_text, file_name):
        """Скачивает файл по ссылке с проверками и обработкой ошибок."""
        try:
            response = requests.get(link, headers=self.headers, timeout=10)  # Устанавливаем таймаут
            response.raise_for_status()  # Генерирует исключение, если статус-код не 200-299

            # Формируем путь сохранения
            save_path = os.path.join(self.file_path, self.main_directory + self.object_name, title_text)

            # Проверяем существование папки, если её нет — создаем
            if not os.path.exists(save_path):
                try:
                    os.makedirs(save_path, exist_ok=True)
                except OSError as e:
                    self.message_signal.emit(f"Ошибка при создании директории {save_path}: {e}")
                    print(f"Ошибка при создании директории {save_path}: {e}")
                    return

            # Пытаемся сохранить файл
            file_path = os.path.join(save_path, file_name)
            with open(file_path, "wb") as file:
                file.write(response.content)
        except requests.exceptions.RequestException as e:
            print(f"Ошибка сети при скачивании {file_name} с {link}: {e}")
        except OSError as e:
            print(f"Ошибка файловой системы при сохранении {file_name}: {e}")
        except Exception as e:
            print(f"Неизвестная ошибка при скачивании {file_name}: {e}")
    
    def get_result_contracts(self, num):
        """Получает результаты контрактов и скачивает документы."""
        try:
            link = f'https://zakupki.gov.ru/epz/order/notice/rpec/documents.html?regNumber={num}0001'
            response = requests.get(link, headers=self.headers)
            soup = BeautifulSoup(response.text, 'lxml')
            
            for block in soup.find_all(class_='col-sm-12 blockInfo'):
                title = block.find(class_='blockInfo__title')
                if title:
                    title_text = title.get_text().strip()
                    
                    for file_block in block.find_all(class_='blockFilesTabDocs'):
                        for link_tag in file_block.find_all('a', href=True, title=True):
                            file_link = link_tag['href']
                            file_name = link_tag['title']
                            
                            if 'download' in file_link or 'file' in file_link:
                                self.downloader(file_link, title_text, file_name)
        except Exception as e:
            self.message_signal.emit(f"Ошибка при получении контрактных документов: {e}")
            print(f"Ошибка при получении контрактных документов: {e}")

    def get_contract_details(self, link):
        self.download_documents(link, 'card-attachments', 'container card-attachments-container', 'title pb-0', 'col-12')

    def get_electronic_documents(self, link):
        self.download_documents(link, 'block-lot card-attachments__block', 'block-lot card-attachments__block', 'title pb-0', 'col-12')
    def sanitize_filename(self,filename, max_length=100):
        # Удаляем все запрещенные символы
        filename = re.sub(r'[\\/:"*?<>|\n\r]+', '', filename)
       
        return filename[:max_length]
    def download_documents(self, link, main_class, container_class, title_class, content_class):
        req = requests.get(url=link, headers=self.headers)
        src = req.text
        soup = BeautifulSoup(src, "lxml")

        containers = soup.find_all(class_=container_class)
        for container in containers:
            title = container.find(class_=title_class)
            if not title:
                continue
            title_text = self.sanitize_filename(title.get_text().strip())

            content_blocks = container.find_all(class_=content_class)
            for block in content_blocks:
                links = block.find_all('a')
                for link in links:
                    try:
                        href = link.get('href')
                        if 'download' in href or 'file' in href:
                            # Используем title, если он есть, иначе fallback на текст ссылки
                            filename = link.get('title', link.get_text()).strip()
                            filename = self.sanitize_filename(filename)

                            self.save_file(href, title_text, filename)

                    except Exception as e:
                        print(f"Ошибка при скачивании файла: {str(e)}")
                        self.message_signal.emit(f"Ошибка при скачивании документов: {str(e)}")
                        continue

    def save_file(self, url, title_text, filename):
        response = requests.get(url, headers=self.headers, stream=True)
        
        if response.status_code == 200:
            # Чистим название файла от лишнего текста в скобках
            filename = re.sub(r'\s*\([^)]*\)', '', filename).strip()

            # Создаём директорию
            directory = f"{self.file_path}/{self.main_directory + self.object_name}/{title_text}"
            os.makedirs(directory, exist_ok=True)

            # Полный путь к файлу
            file_path = os.path.join(directory, filename)

            # Сохраняем файл
            with open(file_path, "wb") as f:
                for chunk in response.iter_content(chunk_size=8192):
                    f.write(chunk)
                    
            print(f"Файл сохранён: {file_path}")
        else:
            self.message_signal.emit(f"Не удалось скачать файл: {url}")
            print(f"Не удалось скачать файл: {url}")
    def status_log(self):
        self.log_data.append(self.status)
        return self.log_data

    # def Make_Json(self):
    #     global_dict = {}
    #     try:
    #         global_dict.update(self.parse_head())
    #         self.status = "Парсинг заголовка завершен"
    #         self.status_log()

    #         global_dict.update(self.main_info_body(self.soup))
    #         self.status = "Основная информация обработана"
    #         self.status_log()

    #         global_dict.update(self.collapse_element(self.soup))
    #         self.status = "Дополнительные элементы обработаны"
    #         self.status_log()

    #         global_dict.update(self.other_info())
    #         self.status = "Дополнительная информация собрана"
    #         self.status_log()

    #         with open(f"{self.file_path}/{self.main_directory}/Все данные закупки №{self.num}.json", "w", encoding="utf-8") as file:
    #             json.dump(global_dict, file, indent=4, ensure_ascii=False)
            
    #         self.status = "Успешная запись файлов"
    #     except Exception as e:
    #         self.status = f"Ошибка записи файлов: {str(e)}"
        
    #     self.status_log()

    def Make_Dock(self, num):
        global_dict = {}
        try:
            global_dict.update(self.parse_head())
            self.status = "Парсинг заголовка завершен"
            self.status_log()

            global_dict.update(self.main_info_body(self.soup))
            self.status = "Основная информация обработана"
            self.status_log()

            global_dict.update(self.collapse_element(self.soup))
            self.status = "Дополнительные элементы обработаны"
            self.status_log()
            
            global_dict.update(self.documents(num))
            global_dict.update(self.other_info())
            
            doc = Document()
            doc.add_heading('Данные о закупке', level=1)
            
            for key, value in global_dict.items():
                doc.add_heading(key, level=2)
                if isinstance(value, list):
                    for item in value:
                        if isinstance(item, dict):
                            for sub_key, sub_value in item.items():
                                doc.add_paragraph(f'{sub_key}: {sub_value}')
                        else:
                            doc.add_paragraph(str(item))
                else:
                    doc.add_paragraph(str(value))
            
            os.makedirs(f"{self.file_path}/{self.main_directory + self.object_name}", exist_ok=True)
            doc_path = f"{self.file_path}/{self.main_directory + self.object_name}/Все данные о закупке №{self.num}.docx"
            doc.save(doc_path)
            self.message_signal.emit(f"Успешная запись файлов №{self.num}")
            self.status = "Успешная запись файлов"
        except Exception as e:
            self.message_signal.emit(f"Ошибка записи файлов: {str(e)}")
            self.status = f"Ошибка записи файлов: {str(e)}"
        
        self.status_log()
        return global_dict

# if __name__ == "__main__":
#     parser = ParserR("TEST 44.csv")
#     parser.make_link_num("0860200000824010096", "C:/Users/Sergey/Download")
#     parser.parse_head()
#     parser.get_supplier_links("0860200000824010096")
#     parser.other_info()
  
